"""Turn a leads CSV into an outreach workbook for the Instagram/Facebook DM team.

python -m leadfinder.outreach_excel results/leads-latest.csv results/AceAds-DM-outreach.xlsx
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

STATUSES = ["Not contacted", "DM sent", "Replied", "Preview sent", "Meeting booked", "Won", "Not interested", "Business closed"]
NAVY, WHITE = "1F3864", "FFFFFF"
PRIORITY_FILL = {"HIGH": "C6EFCE", "VERIFY FIRST": "FFEB9C", "MEDIUM": "DDEBF7"}


def classify(pitch: str) -> tuple[int, str, str]:
    """(sort order, priority, lead type) from the audit finding."""
    if pitch.startswith("no real website"):
        return 1, "HIGH", "No own website"
    if pitch.startswith("no website listed, only social"):
        return 3, "VERIFY FIRST", "No website found"
    if pitch.startswith("security certificate"):
        return 2, "HIGH", "'Not secure' warning"
    if pitch.startswith(("homepage shows an error", "website broken")):
        return 3, "VERIFY FIRST", "Site shows an error"
    return 4, "MEDIUM", "Weak website"


CATEGORY_PHRASE = {
    "barber or hairdresser": "hairdresser", "bar or pub": "bar", "deli or specialty food": "local food store",
    "gift or homewares": "gift shop", "gym or studio": "gym", "massage or spa": "massage", "pet business": "pet shop",
    "allied health": "health practitioner", "trades": "tradie", "takeaway": "takeaway",
}


def opener(row: dict, lead_type: str, platform: str) -> str:
    """First DM, with a slot the team must fill with one real detail from the profile.

    It offers a free preview rather than claiming one exists: build it only when they say yes.
    """
    name = row["name"]
    region = row.get("region") or "your area"
    category = CATEGORY_PHRASE.get(row["category"], row["category"])
    hook = "[one real detail from their page]"
    if lead_type == "No own website":
        problem = (f"noticed that when people in {region} search for a {category} on Google, they only find your "
                   f"{platform}, not a website of your own")
    elif lead_type == "No website found":
        problem = f"couldn't find a website for {name}, only your {platform}"
    elif lead_type == "'Not secure' warning":
        problem = "noticed your website shows a 'Not secure' warning when it's opened on a phone"
    elif lead_type == "Site shows an error":
        problem = "noticed your website was showing an error page when we tried it"
    elif "not mobile-friendly" in row["pitch_angle"]:
        problem = "noticed your website is hard to use on a phone"
    else:
        problem = "think your website could be doing more to bring in customers"
    return (f"Hi {name} team! {hook}. We're AceAds, a web design studio. We {problem}. "
            "We'd be happy to make you a free preview of what a website could look like. Want one? No cost, no obligation.")


def build(csv_path: Path, out_path: Path) -> int:
    rows = list(csv.DictReader(open(csv_path, encoding="utf-8-sig")))
    leads = []
    for r in rows:
        if not (r.get("instagram") or r.get("facebook")):
            continue
        order, priority, lead_type = classify(r["pitch_angle"])
        platform = "Instagram" if r.get("instagram") else "Facebook"
        leads.append((order, r.get("region", ""), r["category"], r["name"], priority, lead_type, platform, r))
    leads.sort(key=lambda x: x[:4])

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    head = ["#", "Priority", "Lead type", "Business", "Category", "Region", "DM on", "Instagram", "Facebook", "Phone",
            "Website", "What we found", "Suggested first DM (fill the [bracket] first)", "Check on Google Maps",
            "Status", "Assigned to", "Date contacted", "Follow-up date", "Notes"]
    ws.append(head)
    for i, (_, _, _, _, priority, lead_type, platform, r) in enumerate(leads, 1):
        ws.append([i, priority, lead_type, r["name"], r["category"], r.get("region", ""), platform, r.get("instagram", ""),
                   r.get("facebook", ""), r["phone"], r["website"], r["pitch_angle"],
                   opener(r, lead_type, platform), r["maps_url"], "Not contacted", "", "", "", ""])
        ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=PRIORITY_FILL[priority])

    for cell in ws[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = [5, 13, 18, 28, 16, 14, 11, 30, 30, 16, 30, 40, 70, 22, 16, 14, 14, 14, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    link_cols = {8: None, 9: None, 11: None, 14: "Open map"}
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for col, label in link_cols.items():
            cell = row[col - 1]
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
                if label:
                    cell.value = label
    last = ws.max_row
    status = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(status)
    status.add(f"O2:O{last}")
    dates = DataValidation(type="date", allow_blank=True)
    ws.add_data_validation(dates)
    dates.add(f"Q2:R{last}")
    for value, colour in [("Won", "A9D08E"), ("Meeting booked", "C6EFCE"), ("Replied", "FFEB9C"),
                          ("Not interested", "D9D9D9"), ("Business closed", "D9D9D9")]:
        ws.conditional_formatting.add(f"A2:S{last}", FormulaRule(formula=[f'$O2="{value}"'],
                                                                  fill=PatternFill("solid", fgColor=colour)))
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:S{last}"

    dash = wb.create_sheet("Tracker")
    dash.append(["Pipeline", "Leads"])
    for s in STATUSES:
        dash.append([s, f'=COUNTIF(Leads!O:O,"{s}")'])
    dash.append([])
    dash.append(["Reply rate (Replied or later / DMs sent)",
                 '=IFERROR((B4+B5+B6+B7+B8)/(B3+B4+B5+B6+B7+B8),0)'])
    dash["B11"].number_format = "0%"
    dash.append(["Leads per priority", ""])
    for p in PRIORITY_FILL:
        dash.append([p, f'=COUNTIF(Leads!B:B,"{p}")'])
    dash.append([])
    dash.append(["Leads per region", ""])
    dash.cell(dash.max_row, 1).font = Font(bold=True)
    for region in sorted({r.get("region", "") for *_, r in leads} - {""}):
        dash.append([region, f'=COUNTIF(Leads!F:F,"{region}")'])
    for cell in dash[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
    dash["A12"].font = Font(bold=True)
    dash.column_dimensions["A"].width = 42
    dash.column_dimensions["B"].width = 12

    rules = wb.create_sheet("How to work this list")
    regions = sorted({r.get("region", "") for *_, r in leads} - {""})
    area = ", ".join(regions) if 0 < len(regions) <= 3 else "Queensland" if regions else "Brisbane"
    lines = [
        (f"AceAds: Instagram / Facebook outreach, {area}", True),
        ("", False),
        ("Before each DM", True),
        ("1. Open 'Check on Google Maps'. Skip the business if it is closed or has almost no reviews.", False),
        ("2. Open their Instagram/Facebook. Replace [one real detail from their page] with something true and specific.", False),
        ("3. 'VERIFY FIRST' rows: Google the business name first. If they DO have a website, check it on your phone and", False),
        ("   rewrite the DM around a problem you can see yourself, or skip them. Never say they have no website unless you checked.", False),
        ("", False),
        ("Sending", True),
        ("4. Send from the AceAds business account, by hand. Max 20 new DMs per account per day.", False),
        ("5. Never paste the identical message twice in a row. Personalise every one; Meta limits accounts that send copy-paste DMs.", False),
        ("6. If they reply yes, ask AceAds for their preview (mockups/ in the repo), then send the screenshot and link.", False),
        ("   Set Status to 'Preview sent'.", False),
        ("7. One polite follow-up after 4-5 days, only if there was no reply. Then stop.", False),
        ("8. If anyone says no or asks us to stop, set 'Not interested' and never contact them again.", False),
        ("", False),
        ("Do not", True),
        ("- Do not add these businesses to bulk email or automated DM tools.", False),
        ("- Do not claim anything about their business you have not checked.", False),
        ("", False),
        ("Source: OpenStreetMap (c) OpenStreetMap contributors, plus links found on the businesses' own websites.", False),
    ]
    for text, bold in lines:
        rules.append([text])
        if bold:
            rules.cell(rules.max_row, 1).font = Font(bold=True, size=13 if rules.max_row == 1 else 11)
    rules.column_dimensions["A"].width = 120

    wb.save(out_path)
    return len(leads)


if __name__ == "__main__":
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "results/leads-latest.csv")
    dst = Path(sys.argv[2] if len(sys.argv) > 2 else "results/AceAds-DM-outreach.xlsx")
    print(f"{build(src, dst)} leads written to {dst}")
