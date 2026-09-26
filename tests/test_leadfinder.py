import datetime as dt

from leadfinder.audit import Audit, analyse_html, audit_website, is_social_only, latest_copyright_year
from leadfinder.places import parse_place
from leadfinder.scoring import demand_score, gap_score, tier

TODAY = dt.date(2026, 9, 25)
GOOD_HTML = """<html><head><title>Smith Plumbing</title>
<meta name="viewport" content="width=device-width">
<meta name="description" content="Brisbane plumbers"></head>
<body><footer>&copy; 2026 Smith Plumbing</footer></body></html>"""
OLD_HTML = "<html><head><title></title></head><body>Copyright 2017 Old Co</body></html>"


def test_modern_site_has_no_issues():
    assert analyse_html(GOOD_HTML, "https://smith.com.au/", 1.2, TODAY) == []


def test_old_site_issues():
    issues = analyse_html(OLD_HTML, "http://old.com.au/", 4.5, TODAY)
    assert issues == [
        "no HTTPS (browser shows 'Not secure')",
        "not mobile-friendly (no viewport tag)",
        "looks outdated (copyright 2017)",
        "missing page title",
        "missing meta description",
        "slow to load (4.5s)",
    ]


def test_copyright_year_takes_latest():
    assert latest_copyright_year("© 2019 - 2024 Co") == 2024
    assert latest_copyright_year("no year here") is None


def test_social_only_and_missing_websites():
    assert is_social_only("https://www.facebook.com/smithplumbing")
    assert not is_social_only("https://smithplumbing.com.au")
    assert audit_website("").status == "none"
    assert audit_website("https://m.facebook.com/x").status == "social_only"


def test_demand_rewards_reviews_and_rating():
    assert demand_score(5.0, 0) == 0
    assert demand_score(5.0, 200) == 40
    assert demand_score(4.8, 150) > demand_score(4.8, 10)
    assert demand_score(4.9, 100) > demand_score(3.5, 100)


def test_gap_score():
    assert gap_score(Audit(status="unreachable")) == 45
    assert gap_score(Audit(status="ok", issues=[])) == 0
    issues = ["no HTTPS (x)", "not mobile-friendly (x)", "looks outdated (x)", "slow to load (x)", "poor mobile performance (x)"]
    assert gap_score(Audit(status="ok", issues=issues)) == 60


def test_tiers():
    assert (tier(75), tier(45), tier(10)) == ("A", "B", "C")


def test_parse_place_skips_closed():
    place = {"id": "p1", "displayName": {"text": "Smith"}, "rating": 4.7, "userRatingCount": 88}
    business = parse_place(place, "plumber")
    assert business.name == "Smith" and business.reviews == 88 and business.website == ""
    assert parse_place({**place, "businessStatus": "CLOSED_PERMANENTLY"}, "plumber") is None


def test_cli_imports_and_requires_key(monkeypatch, tmp_path):
    from leadfinder import cli

    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    assert cli.main(["--source", "google"]) == 1


def test_osm_parse_and_query():
    from leadfinder.osm import build_query, parse_element

    el = {"type": "node", "id": 7, "tags": {"name": "Bean Cafe", "addr:street": "Ann St", "contact:website": "http://bean.au"}}
    b = parse_element(el, "cafe")
    assert b.place_id == "osm:node/7" and b.website == "http://bean.au" and b.address == "Ann St"
    assert "google.com/maps/search" in b.maps_url
    assert parse_element({"type": "node", "id": 8, "tags": {}}, "cafe") is None
    assert build_query(['"amenity"="cafe"'], [-27.7, 152.8, -27.2, 153.25]) == (
        '[out:json][timeout:120];(nwr["amenity"="cafe"](-27.7,152.8,-27.2,153.25););out tags center;'
    )


def test_cli_osm_end_to_end(monkeypatch, tmp_path):
    import csv
    import json

    from leadfinder import cli, osm
    from leadfinder.places import Business

    config = {"area": "Brisbane", "categories": ["cafe", "plumber"], "osm_bbox": [0, 0, 1, 1], "osm_tags": {"cafe": ["x"]}}
    (tmp_path / "config.json").write_text(json.dumps(config))
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.setattr(osm, "search", lambda c, t, b: [
        Business("osm:node/1", "No Site Cafe", c, "", "", "", 0.0, 0, ""),
        Business("osm:node/2", "FB Cafe", c, "", "", "https://facebook.com/fb", 0.0, 0, ""),
    ])
    assert cli.main(["--out", "leads.csv"]) == 0
    raw = (tmp_path / "leads.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM for Excel
    rows = list(csv.DictReader(raw.decode("utf-8-sig").splitlines()))
    assert [(r["name"], r["score"], r["tier"]) for r in rows] == [("FB Cafe", "75", "A"), ("No Site Cafe", "25", "C")]


def test_blocked_site_is_not_reported_broken(monkeypatch):
    import leadfinder.audit as audit_mod

    class Resp:
        status_code = 403

    monkeypatch.setattr(audit_mod.requests, "get", lambda *a, **k: Resp())
    result = audit_website("https://protected.com.au")
    assert result.status == "blocked" and gap_score(result) == 0


def test_osm_fetch_falls_back_to_mirror(monkeypatch):
    from leadfinder import osm

    calls = []

    class Resp:
        def __init__(self, code):
            self.status_code = code

        def raise_for_status(self):
            pass

        def json(self):
            return {"elements": []}

    def fake_post(url, **kwargs):
        calls.append(url)
        return Resp(504 if len(calls) == 1 else 200)

    monkeypatch.setattr(osm.requests, "post", fake_post)
    assert osm.fetch("q") == {"elements": []}
    assert calls == osm.OVERPASS_URLS[:2]


def test_osm_missing_website_is_unknown_not_none():
    assert gap_score(Audit(status="unknown")) < gap_score(Audit(status="none"))


class FakeResp:
    def __init__(self, code, url="https://x.com.au/", text="<html></html>"):
        self.status_code, self.url, self.text = code, url, text


def test_stale_deep_link_falls_back_to_homepage(monkeypatch):
    import leadfinder.audit as audit_mod

    seen = []

    def fake_get(url, **kwargs):
        seen.append(url)
        return FakeResp(404) if url.endswith("/old-page") else FakeResp(200, url, GOOD_HTML)

    monkeypatch.setattr(audit_mod.requests, "get", fake_get)
    result = audit_website("https://cafe.com.au/pages/old-page")
    assert seen == ["https://cafe.com.au/pages/old-page", "https://cafe.com.au/"]
    assert result.status == "ok"


def test_connection_failures_are_classified(monkeypatch):
    import leadfinder.audit as audit_mod

    def raiser(exc):
        def fake_get(*a, **k):
            raise exc
        return fake_get

    monkeypatch.setattr(audit_mod.requests, "get", raiser(audit_mod.requests.exceptions.SSLError()))
    assert audit_website("https://a.com.au").status == "ssl_error"
    monkeypatch.setattr(audit_mod.requests, "get", raiser(audit_mod.requests.ConnectionError()))
    dead = audit_website("https://a.com.au")
    assert dead.status == "dead" and gap_score(dead) < gap_score(Audit(status="social_only"))
    monkeypatch.setattr(audit_mod.requests, "get", lambda *a, **k: FakeResp(500))
    assert audit_website("https://a.com.au").status == "unreachable"


def test_government_sites_excluded():
    from leadfinder.audit import is_excluded

    assert is_excluded("https://indigiscapes.redland.qld.gov.au/info")
    assert not is_excluded("https://cafe.com.au") and not is_excluded("")


def test_bot_challenge_page_is_blocked(monkeypatch):
    import leadfinder.audit as audit_mod

    monkeypatch.setattr(audit_mod.requests, "get", lambda *a, **k: FakeResp(200, "https://x.com.au/", "<html><script>challenge()</script></html>"))
    assert audit_website("https://x.com.au").status == "blocked"


def test_url_without_scheme_and_booking_pages(monkeypatch):
    import leadfinder.audit as audit_mod

    seen = []
    monkeypatch.setattr(audit_mod.requests, "get", lambda url, **k: seen.append(url) or FakeResp(200, url, GOOD_HTML))
    assert audit_website("www.milanigelato.com.au").status == "ok"
    assert seen == ["http://www.milanigelato.com.au"]
    assert audit_website("https://tommytwoblades.gettimely.com/#home").status == "social_only"
    assert audit_website("https://petersandko.wixsite.com/nowhere").status == "social_only"


def test_social_links_and_chains():
    from leadfinder.audit import find_social_links
    from leadfinder.osm import parse_element

    page = ('<a href="https://www.facebook.com/sharer/sharer.php?u=x">share</a>'
            '<a href="https://www.facebook.com/bestcafe/">fb</a><a href="https://instagram.com/bestcafe">ig</a>')
    assert find_social_links(page) == {"facebook": "https://www.facebook.com/bestcafe", "instagram": "https://instagram.com/bestcafe"}

    el = {"type": "node", "id": 1, "tags": {"name": "Bean", "contact:instagram": "@beanbne", "facebook": "https://facebook.com/bean"}}
    b = parse_element(el, "cafe")
    assert b.instagram == "https://www.instagram.com/beanbne" and b.facebook == "https://facebook.com/bean"
    assert parse_element({"type": "node", "id": 2, "tags": {"name": "Starbucks", "brand:wikidata": "Q37158"}}, "cafe") is None


def test_require_social_and_social_listed(monkeypatch, tmp_path):
    import csv
    import json

    from leadfinder import cli, osm
    from leadfinder.places import Business

    config = {"area": "Brisbane", "categories": ["cafe"], "osm_bbox": [0, 0, 1, 1], "osm_tags": {"cafe": ["x"]}}
    (tmp_path / "config.json").write_text(json.dumps(config))
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.setattr(osm, "search", lambda c, t, b: [
        Business("osm:node/1", "Insta Only", c, "", "", "", 0.0, 0, "", instagram="https://www.instagram.com/io"),
        Business("osm:node/2", "FB Site", c, "", "", "https://facebook.com/fbs", 0.0, 0, ""),
        Business("osm:node/3", "Nothing Known", c, "", "", "", 0.0, 0, ""),
    ])
    assert cli.main(["--out", "leads.csv", "--require-social"]) == 0
    rows = list(csv.DictReader(open(tmp_path / "leads.csv", encoding="utf-8-sig")))
    assert [(r["name"], r["tier"], r["facebook"], r["instagram"]) for r in rows] == [
        ("FB Site", "A", "https://facebook.com/fbs", ""),
        ("Insta Only", "B", "", "https://www.instagram.com/io"),
    ]


def test_outreach_workbook(tmp_path):
    import csv

    from openpyxl import load_workbook

    from leadfinder.cli import COLUMNS
    from leadfinder.outreach_excel import build

    base = dict.fromkeys(COLUMNS, "")
    rows = [
        {**base, "name": "Insta Cafe", "category": "cafe", "pitch_angle": "no real website (uses www.instagram.com)",
         "instagram": "https://www.instagram.com/ic", "maps_url": "https://maps.example"},
        {**base, "name": "Old Site Dental", "category": "dentist", "pitch_angle": "security certificate error (x)",
         "facebook": "https://facebook.com/osd"},
        {**base, "name": "No Handle", "category": "cafe", "pitch_angle": "no real website (uses x)"},
    ]
    src = tmp_path / "leads.csv"
    with open(src, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
    assert build(src, tmp_path / "out.xlsx") == 2
    wb = load_workbook(tmp_path / "out.xlsx")
    ws = wb["Leads"]
    assert [ws.cell(r, 4).value for r in (2, 3)] == ["Insta Cafe", "Old Site Dental"]
    assert ws.cell(2, 6).value == "Instagram" and ws.cell(3, 6).value == "Facebook"
    assert "only find your Instagram" in ws.cell(2, 12).value and "[one real detail" in ws.cell(2, 12).value
    assert wb["Tracker"]["A11"].value.startswith("Reply rate")
